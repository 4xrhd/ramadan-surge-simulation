import os
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "src"))
import workbook  # noqa: E402

C = ["rice", "lentils"]


def test_build_workbook_smoke(tmp_path):
    idx = pd.date_range("2020-01-01", periods=4, freq="MS")
    data = pd.DataFrame({"National rice": [40, 41, np.nan, 43.0]}, index=idx)
    cols = pd.MultiIndex.from_tuples([("all", "marginal")] + [(c, f) for c in C for f in ("baseline_m1", "baseline_m2", "flag")])
    wide = pd.DataFrame([[False, 40, 41, 44, 90, 92, 99], [True, 41, 42, 43, np.nan, 90, 95]], index=[2020, 2021], columns=cols)
    draws = pd.DataFrame(np.ones((5, 4)), columns=["rice_b", "rice_s", "lentils_b", "lentils_s"])
    py_risk = pd.DataFrame({"rice": [1.0], "lentils": [2.0], "basket": [3.0]}, index=["Expected price"])
    bt = pd.DataFrame({"year": [2024], "forecast_median": [10.0], "forecast_p5": [8.0], "forecast_p95": [12.0],
                       "actual_price": [11.0], "naive_last_price": [9.0], "actual_surge": [0.05],
                       "surge_p5": [0.0], "surge_p95": [0.1]})
    assumptions = pd.DataFrame({"item": ["Seed"], "value": [42], "note": ["global"]})
    out = workbook.build_workbook(
        tmp_path / "r.xlsx", data=data, flags=pd.Series(1, index=idx[:1]), surge_blocks={"National": wide},
        commodities=C, draws=draws, risk_title="t", python_risk=py_risk,
        scenarios=pd.DataFrame({"scenario": ["Base"], "p95": [1.0]}), backtest=bt,
        assumptions=assumptions, basket={"rice": 40, "lentils": 2})
    wb = load_workbook(out)
    assert wb.sheetnames == ["Data", "Surge", "Risk", "Scenarios", "Backtest", "Assumptions"]
    surge = wb["Surge"]
    assert surge["G3"].value == "=F3/E3-1" and surge["E3"].value == "=AVERAGE(C3:D3)"
    assert surge["L4"].value is None                      # missing lentils baseline skipped
    risk = wb["Risk"]
    assert risk["D3"].value == "=B3*(1+C3)"
    assert "Assumptions!$C$3" in risk["H3"].value
    bt_ws = wb["Backtest"]
    assert bt_ws["J3"].value == "=ABS(B3-E3)/E3"
    assert bt_ws["B5"].value.startswith("=AVERAGE(J3:J3)")
    assert wb["Assumptions"]["C3"].value == 40.0


def test_new_excel_functions_are_prefixed(tmp_path):
    test_build_workbook_smoke(tmp_path)
    wb = load_workbook(tmp_path / "r.xlsx")
    texts = [c.value for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    for fn in ("STDEV.S(", "PERCENTILE.INC("):
        uses = [t for t in texts if fn in t]
        assert uses and all(f"_xlfn.{fn}" in t for t in uses)
